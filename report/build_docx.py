"""Convert report/PreThesis_Report.md to a formatted .docx.

Formatting per the IU-VNU HCMC pre-thesis template:
Times New Roman 13pt, 1.5 line spacing, justified body, A4, margins 2.5 cm
(left 3 cm for binding), page numbers bottom-centre, auto-generated Table of
Contents field, chapter headings on new pages.

The Table of Contents is inserted as a Word TOC *field*: Word populates it on
open (or on Ctrl+A then F9). The List of Figures and List of Tables are kept as
the explicit tables authored in the Markdown, because the figure and table
captions are not SEQ-numbered and a `TOC \\c` field would therefore find nothing.

Usage:
    python report/build_docx.py
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

REPORT_DIR = Path(__file__).resolve().parent
SOURCE = REPORT_DIR / "PreThesis_Report.md"
TARGET = REPORT_DIR / "PreThesis_Report.docx"

BODY_FONT = "Times New Roman"
MONO_FONT = "Consolas"
BODY_SIZE = Pt(13)
LINE_SPACING = 1.5

HEADING_SIZES = {1: Pt(18), 2: Pt(15), 3: Pt(14), 4: Pt(13)}


# ---------------------------------------------------------------------------
# Document setup
# ---------------------------------------------------------------------------
def configure_document(doc: Document) -> None:
    """A4 page, binding margin, base style, heading styles, page-number footer."""
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.right_margin = Cm(2.5)
        section.left_margin = Cm(3.0)  # extra for binding

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    pf = normal.paragraph_format
    pf.line_spacing = LINE_SPACING
    pf.space_after = Pt(6)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for level, size in HEADING_SIZES.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style.font.size = size
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.paragraph_format.line_spacing = LINE_SPACING
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.keep_with_next = True

    add_page_number_footer(doc.sections[0])


def add_page_number_footer(section) -> None:
    """Bottom-centre PAGE field."""
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.font.name = BODY_FONT
    run.font.size = Pt(11)

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, end):
        run._r.append(element)


def add_toc_field(doc: Document) -> None:
    """Insert a TOC field Word populates on open."""
    paragraph = doc.add_paragraph()
    run = paragraph.add_run()

    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and choose “Update Field” to build the Table of Contents."
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, placeholder, end):
        run._r.append(element)


# ---------------------------------------------------------------------------
# Inline formatting
# ---------------------------------------------------------------------------
INLINE = re.compile(r"(\*\*.+?\*\*|\*[^*\n]+?\*|`[^`\n]+?`|\[[^\]]+?\]\([^)]+?\))")


def add_inline(paragraph, text: str, base_size: Pt | None = None) -> None:
    """Render **bold**, *italic*, `code`, and [links](x) into runs."""
    text = text.replace("$$", "").replace("\\", "")
    for token in INLINE.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = MONO_FONT
            run.font.size = Pt(11)
        elif token.startswith("["):
            label = re.match(r"\[([^\]]+)\]", token)
            run = paragraph.add_run(label.group(1) if label else token)
        else:
            run = paragraph.add_run(token)
        if base_size is not None and run.font.size is None:
            run.font.size = base_size


# ---------------------------------------------------------------------------
# Block builders
# ---------------------------------------------------------------------------
def add_table(doc: Document, rows: list[str]) -> None:
    """Render a Markdown pipe table."""
    parsed = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    parsed = [r for i, r in enumerate(parsed) if not (i == 1 and set("".join(r)) <= set("-: "))]
    if not parsed:
        return
    width = max(len(r) for r in parsed)
    parsed = [r + [""] * (width - len(r)) for r in parsed]

    table = doc.add_table(rows=len(parsed), cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for i, row in enumerate(parsed):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(paragraph, cell_text, base_size=Pt(11))
            for run in paragraph.runs:
                run.font.size = Pt(11)
                if i == 0:
                    run.bold = True
    doc.add_paragraph()


def add_code_block(doc: Document, lines: list[str], language: str) -> None:
    """Monospaced, unjustified block. Mermaid sources get a leading note."""
    if language == "mermaid":
        note = doc.add_paragraph()
        note.paragraph_format.space_after = Pt(2)
        run = note.add_run("[Diagram — Mermaid source; render to an image before submission]")
        run.italic = True
        run.font.size = Pt(11)
    for line in lines:
        paragraph = doc.add_paragraph()
        pf = paragraph.paragraph_format
        pf.line_spacing = 1.0
        pf.space_after = Pt(0)
        pf.left_indent = Cm(0.6)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = paragraph.add_run(line if line.strip() else " ")
        run.font.name = MONO_FONT
        run.font.size = Pt(10)
    doc.add_paragraph()


# ---------------------------------------------------------------------------
# Main conversion
# ---------------------------------------------------------------------------
def convert(source: Path, target: Path) -> dict:
    lines = source.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)

    stats = {"headings": 0, "tables": 0, "code_blocks": 0, "paragraphs": 0, "page_breaks": 0}
    in_title_page = True
    i = 0

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # Explicit page break
        if stripped == "\\newpage":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            stats["page_breaks"] += 1
            in_title_page = False
            i += 1
            continue

        # Fenced code block
        if stripped.startswith("```"):
            language = stripped[3:].strip().lower()
            block: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            i += 1
            add_code_block(doc, block, language)
            stats["code_blocks"] += 1
            continue

        # Table
        if stripped.startswith("|") and stripped.endswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            add_table(doc, block)
            stats["tables"] += 1
            continue

        # Heading
        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2)
            paragraph = doc.add_heading("", level=min(level, 4))
            add_inline(paragraph, text)
            if in_title_page:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            stats["headings"] += 1
            i += 1

            # Replace the Table of Contents placeholder with a real TOC field.
            if text.strip().upper().startswith("TABLE OF CONTENTS"):
                add_toc_field(doc)
                while i < len(lines) and not re.match(r"^##\s", lines[i].strip()):
                    i += 1
            continue

        # Horizontal rule
        if stripped in {"---", "***", "___"}:
            i += 1
            continue

        # Blank
        if not stripped:
            i += 1
            continue

        # HTML line-break markers used on the title page
        if stripped in {"<br>", "<br><br>"}:
            doc.add_paragraph()
            i += 1
            continue

        # Blockquote
        if stripped.startswith(">"):
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(1.0)
            add_inline(paragraph, stripped.lstrip("> ").strip())
            for run in paragraph.runs:
                run.italic = True
            stats["paragraphs"] += 1
            i += 1
            continue

        # List item
        bullet = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if bullet:
            indent = len(bullet.group(1)) // 2
            style = "List Bullet" if bullet.group(2) in {"-", "*"} else "List Number"
            paragraph = doc.add_paragraph(style=style)
            paragraph.paragraph_format.left_indent = Cm(0.8 + 0.6 * indent)
            paragraph.paragraph_format.line_spacing = LINE_SPACING
            add_inline(paragraph, bullet.group(3))
            stats["paragraphs"] += 1
            i += 1
            continue

        # Body paragraph (join wrapped lines)
        block = [stripped]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("#", "|", ">", "```", "---", "<br>", "\\newpage"))
                    or re.match(r"^(\s*)([-*]|\d+\.)\s+", lines[i])):
                break
            block.append(nxt)
            i += 1

        paragraph = doc.add_paragraph()
        if in_title_page:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(paragraph, " ".join(block))
        stats["paragraphs"] += 1

    doc.save(target)
    return stats


if __name__ == "__main__":
    result = convert(SOURCE, TARGET)
    print(f"Wrote {TARGET}")
    for key, value in result.items():
        print(f"  {key:14s} {value}")
